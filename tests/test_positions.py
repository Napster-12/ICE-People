"""Phase 4: approved positions register, ad-hoc positions, position status, changes, scoping and the vacancy
dashboard (spec §15.3)."""
import io
from datetime import timedelta
from decimal import Decimal

import pytest

from app.extensions import db
from app.models import ApprovedPosition, AuditLog, Department, Employee, JobGrade, Notification, PositionRegister
from app.services import positions, workflow
from app.services.dates import today
from app.services.errors import DomainError
from app.services.positions import FLOW
from tests.conftest import financial_year, login, make_position, make_requisition, user


def approve(entity_type, entity_id, *actors):
    inst = workflow.latest_instance(entity_type, entity_id)
    for actor in actors:
        workflow.act(inst.current_step, user(actor), "approve", f"ok from {actor}")
    return inst


def values(**overrides):
    sal = Department.query.filter_by(code="SAL").first()
    v = {"title": "Branch Administrator", "department_id": sal.id, "cost_centre_id": sal.cost_centres[0].id,
         "job_grade_id": JobGrade.query.filter_by(name="Senior").first().id, "budgeted_ctc": Decimal("300000"),
         "minimum_requirements": "Matric\n2 years' admin experience", "critical": False}
    v.update(overrides)
    return v


def next_fy():
    return financial_year() + 1


# ------------------------------------------------------------- register

def test_register_load_and_approval(app):
    fy = next_fy()
    a = positions.create_planned(values(), user("hradmin"), fy)
    b = positions.create_planned(values(title="Team Leader"), user("hradmin"), fy)
    assert a.status == "Planned" and a.position_number.startswith("POS-")
    holder = user("employee").employee
    holder.approved_position_id = b.id
    inst = positions.submit_register(fy, user("hradmin"))
    assert [s.name for s in inst.steps] == ["Finance", "Group HR Manager"]
    with pytest.raises(DomainError, match="approved or in approval"):
        positions.create_planned(values(), user("hradmin"), fy)
    approve("position_register", inst.entity_id, "finance", "grouphr")
    assert PositionRegister.query.filter_by(financial_year=fy).one().status == "Approved"
    assert a.status == "Vacant" and a.vacant_since == today()
    assert b.status == "Filled"


def test_returned_register_can_be_changed_and_resent(app):
    fy = next_fy()
    p = positions.create_planned(values(), user("hradmin"), fy)
    inst = positions.submit_register(fy, user("hradmin"))
    workflow.act(inst.current_step, user("finance"), "return", "Cut one position")
    assert positions.register_for(fy).status == "Returned"
    positions.delete_planned(p, user("hradmin"))
    with pytest.raises(DomainError, match="Add positions"):
        positions.submit_register(fy, user("hradmin"))
    positions.create_planned(values(title="Clerk"), user("hradmin"), fy)
    assert positions.submit_register(fy, user("hradmin")).round == 2


def test_after_approval_hr_changes_grade_and_ctc_but_not_department(app):
    p = make_position()
    sal_cc, fin = p.cost_centre_id, Department.query.filter_by(code="FIN").first()
    manager = JobGrade.query.filter_by(name="Manager").first()
    positions.update(p, {"title": "Senior Consultant", "job_grade_id": manager.id, "budgeted_ctc": Decimal("400000"),
                         "department_id": fin.id, "cost_centre_id": fin.cost_centres[0].id}, user("hradmin"))
    assert p.title == "Senior Consultant" and p.job_grade.name == "Manager" and p.budgeted_ctc == Decimal("400000")
    assert p.department.code == "SAL" and p.cost_centre_id == sal_cc               # fixed once approved
    assert AuditLog.query.filter_by(entity="position", action="update").count() == 1


def test_hr_completes_grade_and_ctc_on_ad_hoc_in_approval(app):
    inst = positions.request_ad_hoc(values(justification="x" * 60, budgeted_ctc=None), user("linemanager"))
    p = db.session.get(ApprovedPosition, inst.entity_id)
    senior = JobGrade.query.filter_by(name="Senior").first()
    positions.update(p, {"title": p.title, "job_grade_id": senior.id, "budgeted_ctc": Decimal("275000")}, user("hrbp"))
    assert p.status == "Pending approval" and p.job_grade_id == senior.id and p.budgeted_ctc == Decimal("275000")


def test_validation(app):
    fin = Department.query.filter_by(code="FIN").first()
    with pytest.raises(DomainError, match="belong to the chosen department"):
        positions.create_planned(values(cost_centre_id=fin.cost_centres[0].id), user("hradmin"), next_fy())
    with pytest.raises(DomainError, match="amount in rand"):
        positions.create_planned(values(budgeted_ctc=None, budgeted_ctc_raw="lots"), user("hradmin"), next_fy())
    with pytest.raises(DomainError, match="position title"):
        positions.create_planned(values(title=" "), user("hradmin"), next_fy())
    with pytest.raises(DomainError, match="job grade"):                 # e.g. a form loaded before grades changed
        positions.create_planned(values(job_grade_id=9999), user("hradmin"), next_fy())


def test_status_machine():
    assert FLOW.can("Vacant", "In recruitment")
    with pytest.raises(DomainError):
        FLOW.check("Planned", "In recruitment")
    with pytest.raises(DomainError):
        FLOW.check("Abolished", "Vacant")


# ------------------------------------------------------------- CSV load

def test_csv_import_is_all_or_nothing(app):
    fy = f"{next_fy()}/{str(next_fy() + 1)[-2:]}"
    header = ",".join(positions.CSV_COLUMNS)
    good = f"{fy},SAL,SAL-100,Consultant,junior,,,250000,,Matric | Sales experience"
    bad = f"{fy},SAL,FIN-100,Clerk,4,NOPE,,abc,POS-9,"
    saved, errors = positions.import_csv("\n".join([header, good, bad]), user("hradmin"))
    assert saved == 0 and len(errors) == 1
    assert all(word in errors[0] for word in ("doesn't belong", "unknown grade", "unknown branch",
                                              "isn't a valid number", "unknown reports-to"))
    saved, errors = positions.import_csv("\n".join([header, good, good]), user("hradmin"))
    assert saved == 2 and not errors
    p = ApprovedPosition.query.filter_by(financial_year=next_fy()).first()
    assert p.minimum_requirements == "Matric\nSales experience" and p.status == "Planned"
    assert positions.import_csv("a,b\n1,2", user("hradmin"))[1][0].startswith("Row 1")


# ------------------------------------------------------------- ad hoc

def test_ad_hoc_request_approved_joins_register(app):
    inst = positions.request_ad_hoc(values(justification="x" * 60), user("linemanager"))
    p = db.session.get(ApprovedPosition, inst.entity_id)
    assert p.status == "Pending approval" and p.is_ad_hoc
    assert [s.name for s in inst.steps] == ["Finance (unbudgeted)", "GCOO / Exco"]          # raised by a manager
    with pytest.raises(Exception):
        make_requisition(user("linemanager"), approved_position_id=p.id)
    db.session.rollback()


def test_ad_hoc_full_approval_then_requisition(app):
    inst = positions.request_ad_hoc(values(justification="x" * 60), user("linemanager"))
    approve("position", inst.entity_id, "finance", "gcoo")
    p = db.session.get(ApprovedPosition, inst.entity_id)
    assert p.status == "Vacant"
    note = Notification.query.filter_by(user_id=user("linemanager").id).order_by(Notification.id.desc()).first()
    assert "raise a requisition" in note.body
    req = make_requisition(user("linemanager"), approved_position_id=p.id)
    assert req.status == "In Approval" and req.title == "Branch Administrator"


def test_ad_hoc_return_resubmit_and_reject(app):
    with pytest.raises(DomainError, match="at least 50"):
        positions.request_ad_hoc(values(justification="short"), user("linemanager"))
    inst = positions.request_ad_hoc(values(justification="x" * 60), user("linemanager"))
    p = db.session.get(ApprovedPosition, inst.entity_id)
    workflow.act(inst.current_step, user("finance"), "return", "Which branch?")
    assert p.status == "Returned"
    with pytest.raises(DomainError, match="Only the requester"):
        positions.request_ad_hoc(values(justification="y" * 60), user("hod"), p)
    inst2 = positions.request_ad_hoc(values(justification="y" * 60, title="Durban Administrator"),
                                     user("linemanager"), p)
    assert inst2.round == 2 and p.status == "Pending approval" and p.title == "Durban Administrator"
    workflow.act(inst2.current_step, user("finance"), "reject", "Not needed")
    assert p.status == "Rejected"


def test_ad_hoc_raised_by_hr_starts_with_department_manager(app):
    inst = positions.request_ad_hoc(values(justification="x" * 60), user("hrmanager"))
    assert [s.name for s in inst.steps] == ["Department manager", "Finance (unbudgeted)", "GCOO / Exco"]
    assert inst.current_step.approver_id == user("hod").id                # head of Sales, the position's department
    approve("position", inst.entity_id, "hod", "finance", "gcoo")
    assert db.session.get(ApprovedPosition, inst.entity_id).status == "Vacant"


# ------------------------------------------------- status from requisitions and employees

def test_requisition_approval_and_close_move_position(app):
    from app.services import requisitions
    req = make_requisition(user("hod"))
    position = req.approved_position
    assert position.status == "Vacant"
    approve("requisition", req.id, "linemanager", "finance", "grouphr", "gcoo")
    assert position.status == "In recruitment"
    requisitions.close_unfilled(req, user("hrbp"), "Budget cut")
    assert position.status == "Vacant"


def test_linking_and_unlinking_an_employee_fills_and_vacates(app):
    p = make_position()
    p.vacant_since = today() - timedelta(days=20)
    emp = user("employee").employee
    emp.approved_position_id = p.id
    positions.employee_changed(emp)
    assert p.status == "Filled" and p.days_to_fill == 20 and p.vacant_since is None
    emp.status = "Exited"
    positions.employee_changed(emp)
    assert p.status == "Vacant" and p.last_holder_id == emp.id and p.vacant_since == today()
    req = make_requisition(user("linemanager"), approved_position_id=p.id)
    assert req.headcount_type == "Replacement" and req.replaced_employee_id == emp.id


# ------------------------------------------------- freeze / unfreeze / abolish / regrade

def test_freeze_unfreeze_abolish_through_approval(app):
    p = make_position()
    inst = positions.request_change(p, "Freeze", "Hiring pause", user("hradmin"))
    assert [s.name for s in inst.steps] == ["Finance", "Group HR Manager"]
    with pytest.raises(DomainError, match="already has a change"):
        positions.request_change(p, "Abolish", "x", user("hradmin"))
    approve("position_change", inst.entity_id, "finance", "grouphr")
    assert p.status == "Frozen"
    with pytest.raises(Exception, match="not vacant"):
        make_requisition(user("linemanager"), approved_position_id=p.id)
    db.session.rollback()


def test_unfreeze_abolish_and_regrade(app):
    p = make_position(status="Frozen")
    approve("position_change", positions.request_change(p, "Unfreeze", "Back on", user("hradmin")).entity_id,
            "finance", "grouphr")
    assert p.status == "Vacant"
    senior = JobGrade.query.filter_by(name="Senior").first()
    with pytest.raises(DomainError, match="new grade"):
        positions.request_change(p, "Regrade", "Bigger role", user("hradmin"), p.job_grade_id)
    approve("position_change", positions.request_change(p, "Regrade", "Bigger role", user("hradmin"),
                                                        senior.id).entity_id, "finance", "grouphr")
    assert p.job_grade_id == senior.id
    approve("position_change", positions.request_change(p, "Abolish", "Restructure", user("hradmin")).entity_id,
            "finance", "grouphr")
    assert p.status == "Abolished"
    with pytest.raises(DomainError, match="can't be frozen"):
        positions.request_change(p, "Freeze", "x", user("hradmin"))
    with pytest.raises(DomainError, match="Give a reason"):
        positions.request_change(make_position(), "Freeze", " ", user("hradmin"))


def test_rejected_change_leaves_position_alone(app):
    p = make_position()
    inst = positions.request_change(p, "Abolish", "Restructure", user("hradmin"))
    workflow.act(inst.current_step, user("finance"), "reject", "Keep it")
    assert p.status == "Vacant" and positions.pending_change(p) is None


# ------------------------------------------------------------- scoping

def test_scoping(app):
    lm_emp = user("linemanager").employee
    own = make_position(title="Sales Team Leader")
    lm_emp.approved_position_id = own.id
    reporting = make_position(title="Reports to the team leader", reports_to_id=own.id)
    elsewhere = make_position(title="Reports to someone else", reports_to_id=make_position().id)
    hr = Department.query.filter_by(code="HR").first()
    hr_position = make_position(department_id=hr.id, cost_centre_id=hr.cost_centres[0].id)
    lm = user("linemanager")
    assert positions.can_view(lm, reporting)
    assert not positions.can_view(lm, elsewhere)
    assert not positions.can_view(lm, hr_position)
    assert positions.can_view(user("hod"), elsewhere)          # the Department Manager sees the whole department
    assert not positions.can_view(user("otherhod"), reporting)
    assert positions.can_view(user("finance"), hr_position)
    assert positions.can_view(user("exco"), hr_position)
    assert reporting in positions.raisable_query(lm).all()


# ------------------------------------------------------------- dashboard

def test_vacancy_dashboard_counts(app):
    from app.services import requisitions
    make_position(status="Filled")
    make_position(status="Frozen")
    advertised = make_requisition(user("hod"))
    approve("requisition", advertised.id, "linemanager", "finance", "grouphr", "gcoo")
    requisitions.set_status(advertised, "Advertised")
    approved_only = make_requisition(user("hod"))
    approve("requisition", approved_only.id, "linemanager", "finance", "grouphr", "gcoo")
    make_position(status="Planned")                          # not on the register yet
    data = positions.vacancy_dashboard(user("hradmin"), financial_year())
    t = data["totals"]
    # conftest's 2 vacant + filled + frozen + 2 from the requisitions
    assert (t["approved"], t["filled"], t["frozen"], t["in_recruitment"], t["not_advertised"]) == (6, 1, 1, 1, 3)
    assert t["vacant"] == 5 and t["rate"] == 83
    assert data["rows"][0]["approved"] == 6
    assert positions.vacancy_dashboard(user("hradmin"), financial_year(), "branch")["rows"][0]["name"] == "Not set"


def test_routes(client):
    p = make_position()
    db.session.commit()
    login(client, "employee")
    assert client.get("/positions/").status_code == 403
    assert client.get("/positions/dashboard").status_code == 403
    client.post("/logout")
    login(client, "hradmin")
    assert p.position_number.encode() in client.get("/positions/").data
    assert client.get(f"/positions/{p.id}").status_code == 200
    assert client.get(f"/positions/{p.id}/edit").status_code == 200
    page = client.get("/positions/dashboard?group_by=department")
    assert page.status_code == 200 and b"Vacancies by department" in page.data
    x = client.get("/positions/dashboard?format=xlsx")
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(x.data))
    assert wb.sheetnames == ["Vacancies", "Positions"] and wb["Positions"].max_row == 4
    assert client.get("/positions/import/template.csv").data.startswith(b"financial_year,")
    r = client.post("/positions/new", data={"fy": next_fy(), **{k: v for k, v in values().items() if k != "critical"}})
    assert r.status_code == 302
    r = client.post(f"/positions/{p.id}/change", data={"kind": "Freeze", "reason": "Pause"})
    assert r.status_code == 302 and positions.pending_change(p) is not None
    assert client.get("/admin/org-structure").status_code == 200
    client.post("/admin/org-structure", data={"kind": "region", "code": "kzn", "name": "KwaZulu-Natal"})
    assert AuditLog.query.filter_by(action="config.region").count() == 1


def test_ad_hoc_routes(client):
    login(client, "linemanager")
    assert client.get("/positions/ad-hoc/new").status_code == 200
    r = client.post("/positions/ad-hoc/new", data={**{k: v for k, v in values().items() if k != "critical"},
                                                   "justification": "z" * 60})
    assert r.status_code == 302
    p = ApprovedPosition.query.filter_by(is_ad_hoc=True).one()
    assert p.status == "Pending approval" and p.requested_by_id == user("linemanager").id
    assert p.position_number.encode() in client.get("/positions/ad-hoc").data
    assert client.get(f"/positions/{p.id}").status_code == 200


def test_department_manager_ad_hoc_prefilled_and_locked(client):
    from app.models import JobFamily
    sal, fin = (Department.query.filter_by(code=c).first() for c in ("SAL", "FIN"))
    family = JobFamily(code="SLS", name="Sales and Distribution")
    db.session.add(family)
    db.session.flush()
    sal.job_family_id = family.id
    db.session.commit()
    hod = user("hod")                                                  # Department Manager of Sales
    own = make_position(title="Head of Sales")                         # the Department Manager's own position
    hod.employee.approved_position_id = own.id
    db.session.commit()
    assert positions.ad_hoc_defaults(hod) == {"department_id": sal.id, "cost_centre_id": hod.employee.cost_centre_id,
                                              "job_family_id": family.id, "reports_to_id": own.id}
    assert positions.ad_hoc_defaults(user("hradmin")) == {}            # HR chooses any department
    with pytest.raises(DomainError, match="only request ad-hoc positions for Sales"):
        positions.request_ad_hoc(values(justification="x" * 60, department_id=fin.id,
                                        cost_centre_id=fin.cost_centres[0].id), hod)
    db.session.rollback()
    other = make_position(title="Somewhere else")
    inst = positions.request_ad_hoc(values(justification="x" * 60, reports_to_id=other.id), hod)
    assert db.session.get(ApprovedPosition, inst.entity_id).reports_to_id == own.id   # forced to the manager
    db.session.rollback()

    login(client, "hod")
    page = client.get("/positions/ad-hoc/new").get_data(as_text=True)
    assert '<select id="reports_to_id" disabled>' in page and own.position_number in page
    assert '<select id="reports_to_id" name="reports_to_id"' not in page
    assert f'<input type="hidden" name="department_id" value="{sal.id}">' in page
    assert f'<option value="{family.id}" selected>' in page
    assert f'<option value="{hod.employee.cost_centre_id}" data-dept="{sal.id}" selected>' in page


def test_reports_to_defaults_to_the_department_manager(app):
    hod = user("hod")                                                  # Department Manager (head) of Sales
    head = make_position(title="Head of Sales", status="Filled")
    hod.employee.approved_position_id = head.id
    other = make_position(title="Somebody else")
    db.session.commit()
    assert positions.department_manager_position(hod.employee.department_id) == head
    planned = positions.create_planned(values(), user("hradmin"), next_fy())
    assert planned.reports_to_id == head.id                            # empty: filled in
    chosen = positions.create_planned(values(reports_to_id=other.id), user("hradmin"), next_fy())
    assert chosen.reports_to_id == other.id                            # HR's choice is kept
    inst = positions.request_ad_hoc(values(justification="x" * 60), user("hradmin"))
    assert db.session.get(ApprovedPosition, inst.entity_id).reports_to_id == head.id
    assert positions.reference_data()["manager_positions"][hod.employee.department_id] == head.id


def test_hr_ad_hoc_department_not_locked(client):
    login(client, "hradmin")
    page = client.get("/positions/ad-hoc/new").get_data(as_text=True)
    assert 'name="department_id" value=' not in page and '<select id="department_id" name="department_id"' in page


def test_requisition_form_shows_job_family(client):
    login(client, "hod")
    assert b"Job family" in client.get("/requisitions/new").data


def test_employee_form_fills_position(client):
    p = make_position()
    emp = user("employee").employee
    db.session.commit()
    login(client, "hradmin")
    client.post(f"/admin/employees/{emp.id}", data={
        "employee_number": emp.employee_number, "first_name": emp.first_name, "last_name": emp.last_name,
        "department_id": emp.department_id, "approved_position_id": p.id, "status": "Permanent",
        "employment_type": "Permanent"})
    assert db.session.get(ApprovedPosition, p.id).status == "Filled"
    assert db.session.get(Employee, emp.id).approved_position_id == p.id
