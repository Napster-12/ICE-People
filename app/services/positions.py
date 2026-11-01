"""Approved positions register (spec §15.3): yearly register and its approval, ad-hoc position requests,
automatic position status, freeze / unfreeze / abolish / regrade through approval, CSV load, scoping and the
vacancy dashboard."""
import csv
import io
import math
from datetime import date
from decimal import Decimal, InvalidOperation

from sqlalchemy import or_

from app.extensions import db
from app.models import (ApprovedPosition, Branch, CostCentre, Department, Employee, JobFamily, JobGrade,
                        PositionChange, PositionRegister, Requisition)
from app.models.user import COMPLIANCE, EXCO, FINANCE, GCOO, GROUP_HR, HOD, HR_ADMIN, HR_MANAGER, HR_ROLES, HRBP, \
    LINE_MANAGER
from app.services import audit, budget, workflow
from app.services.dates import now, today
from app.services.errors import DomainError, StateMachine
from app.services.notifications import notify
from app.services.refs import next_reference

FLOW = StateMachine("position", {
    "Planned": {"Vacant", "Filled"},
    "Pending approval": {"Vacant", "Returned", "Rejected"},
    "Returned": {"Pending approval"},
    "Rejected": set(),
    "Vacant": {"In recruitment", "Filled", "Frozen", "Abolished"},
    "In recruitment": {"Vacant", "Filled"},
    "Filled": {"Vacant"},
    "Frozen": {"Vacant", "Abolished"},
    "Abolished": set(),
})

EDIT_ROLES = {HR_ADMIN, GROUP_HR, HR_MANAGER, HRBP}           # load and maintain the register
SEE_ALL_ROLES = HR_ROLES | {FINANCE, GCOO, EXCO, COMPLIANCE}
REQUEST_ROLES = {LINE_MANAGER, HOD} | HR_ROLES                # may request an ad-hoc position
VIEW_ROLES = SEE_ALL_ROLES | {HOD, LINE_MANAGER}

ON_REGISTER = ("Vacant", "In recruitment", "Filled", "Frozen")
VACANT = ("Vacant", "In recruitment", "Frozen")
OPEN_REQUISITION = ("Draft", "Submitted", "In Approval", "Returned", "Approved", "Advertised", "Interviewing", "Offer")
ADVERTISED_REQUISITION = ("Advertised", "Interviewing", "Offer")
ALWAYS_EDITABLE = ["title", "branch_id", "job_family_id", "reports_to_id", "minimum_requirements", "critical"]
PLANNED_EDITABLE = ALWAYS_EDITABLE + ["department_id", "cost_centre_id", "job_grade_id", "budgeted_ctc"]
# HR may change the grade and budgeted CTC at any stage, an ad-hoc request in approval included (31 Oct 2026).
HR_EDITABLE = ALWAYS_EDITABLE + ["job_grade_id", "budgeted_ctc"]
CHANGE_FROM = {"Freeze": ("Vacant",), "Unfreeze": ("Frozen",), "Abolish": ("Vacant", "Frozen"),
               "Regrade": ("Vacant", "Filled", "Frozen")}
MIN_JUSTIFICATION = 50


def snapshot(p):
    out = {}
    for f in PLANNED_EDITABLE + ["status", "financial_year", "is_ad_hoc"]:
        v = getattr(p, f)
        out[f] = v if v is None or isinstance(v, (bool, int, str)) else str(v)
    return out


# ------------------------------------------------------------------ status

def set_status(position, new, detail=None):
    FLOW.check(position.status, new)
    old = position.status
    position.status = new
    if new == "Vacant" and old != "In recruitment":
        position.vacant_since = position.vacant_since if old in ("Planned", "Frozen") and position.vacant_since \
            else today()
    if new == "Filled":
        position.filled_on = today()
        position.days_to_fill = (today() - position.vacant_since).days if position.vacant_since else None
        position.vacant_since = None
    audit.record("status_change", "position", position.id, before={"status": old}, after={"status": new},
                 detail=detail)


def employee_changed(emp, old_position_id=None):
    """Keep position status in step with the employee linked to it (filled / vacated)."""
    for pid in {old_position_id, emp.approved_position_id} - {None}:
        position = db.session.get(ApprovedPosition, pid)
        db.session.refresh(position, ["holders"])
        current = position.holder
        if current is not None and position.status in ("Vacant", "In recruitment"):
            set_status(position, "Filled", detail=f"{current.full_name} linked to the position")
        elif current is None and position.status == "Filled":
            position.last_holder_id = emp.id
            set_status(position, "Vacant", detail=f"{emp.full_name} no longer holds the position")


def open_requisition(position, exclude=None):
    q = Requisition.query.filter(Requisition.approved_position_id == position.id,
                                 Requisition.status.in_(OPEN_REQUISITION))
    if exclude is not None and exclude.id is not None:
        q = q.filter(Requisition.id != exclude.id)
    return q.order_by(Requisition.id.desc()).first()


def start_recruitment(position, req):
    if position is not None and position.status == "Vacant":
        set_status(position, "In recruitment", detail=f"{req.reference} approved")


def stop_recruitment(position, req):
    if position is not None and position.status == "In recruitment":
        set_status(position, "Vacant", detail=f"{req.reference} closed")


def requisition_problems(req):
    """Why a requisition can't be submitted against its position (empty list if it can)."""
    position = db.session.get(ApprovedPosition, req.approved_position_id) if req.approved_position_id else None
    if position is None:
        return ["Choose a vacant approved position. If the position isn't on the register, request an ad-hoc "
                "position first."]
    if position.status != "Vacant":
        return [f"{position.position_number} is {position.status.lower()}, not vacant."]
    other = open_requisition(position, exclude=req)
    if other is not None:
        return [f"{position.position_number} already has an open requisition ({other.reference})."]
    return []


# ----------------------------------------------------------------- scoping

def own_position_ids(user):
    emp = user.employee
    return [emp.approved_position_id] if emp and emp.approved_position_id else []


def visible_query(user):
    q = ApprovedPosition.query
    if user.has_role(*SEE_ALL_ROLES):
        return q
    conditions = [ApprovedPosition.requested_by_id == user.id]
    if user.has_role(HOD) and user.department_id:
        conditions.append(ApprovedPosition.department_id == user.department_id)
    if user.has_role(LINE_MANAGER):
        mine = own_position_ids(user)
        if mine:
            conditions.append(ApprovedPosition.reports_to_id.in_(mine))
        if user.department_id:
            conditions.append(db.and_(ApprovedPosition.department_id == user.department_id,
                                      ApprovedPosition.reports_to_id.is_(None)))
    return q.filter(or_(*conditions))


def can_view(user, position):
    return visible_query(user).filter(ApprovedPosition.id == position.id).first() is not None


def raisable_query(user, keep=None):
    """Vacant positions this user may raise a requisition for (plus `keep`, the one already chosen)."""
    busy = (db.session.query(Requisition.approved_position_id)
            .filter(Requisition.approved_position_id.isnot(None), Requisition.status.in_(OPEN_REQUISITION)))
    if keep is not None and keep.id is not None:
        busy = busy.filter(Requisition.id != keep.id)
    q = visible_query(user).filter(ApprovedPosition.status == "Vacant", ApprovedPosition.id.notin_(busy))
    return q.order_by(ApprovedPosition.title, ApprovedPosition.position_number)


# ---------------------------------------------------------------- register

def register_for(financial_year, create=False):
    reg = PositionRegister.query.filter_by(financial_year=financial_year).first()
    if reg is None and create:
        reg = PositionRegister(financial_year=financial_year, status="Draft")
        db.session.add(reg)
        db.session.flush()
    return reg


def register_editable(financial_year):
    reg = register_for(financial_year)
    return reg is None or reg.status in ("Draft", "Returned", "Rejected")


def _clean(values):
    v = dict(values)
    v["title"] = (v.get("title") or "").strip()
    v["minimum_requirements"] = (v.get("minimum_requirements") or "").strip() or None
    return v


def department_manager_position(department_id, exclude=None):
    """The position held by a department's Department Manager (its head): the default "reports to" for the
    department's positions. None if the department has no manager or they hold no position on the register."""
    dept = db.session.get(Department, department_id or 0)
    position = dept.hod.approved_position if dept is not None and dept.hod is not None else None
    return None if position is None or (exclude is not None and position.id == exclude.id) else position


def _default_reports_to(values, position=None):
    """Fill an empty "reports to" with the Department Manager's position."""
    if not values.get("reports_to_id"):
        manager = department_manager_position(values.get("department_id"), exclude=position)
        values["reports_to_id"] = manager.id if manager else None
    return values


def validate(values, position=None):
    errors = []
    if not values.get("title"):
        errors.append("Enter a position title.")
    dept = db.session.get(Department, values.get("department_id") or 0)
    if dept is None:
        errors.append("Choose a department.")
    cc = db.session.get(CostCentre, values.get("cost_centre_id") or 0)
    if cc is None:
        errors.append("Choose a cost centre.")
    elif dept is not None and cc.department_id != dept.id:
        errors.append("The cost centre must belong to the chosen department.")
    if db.session.get(JobGrade, values.get("job_grade_id") or 0) is None:     # also catches a stale form's old grade
        errors.append("Choose a job grade.")
    ctc = values.get("budgeted_ctc")
    if values.get("budgeted_ctc_raw") and ctc is None:
        errors.append("Enter the budgeted CTC as an amount in rand, e.g. 250000.00.")
    elif ctc is not None and ctc < 0:
        errors.append("The budgeted CTC can't be negative.")
    reports_to = values.get("reports_to_id")
    if reports_to and position is not None and reports_to == position.id:
        errors.append("A position can't report to itself.")
    return errors


def create_planned(values, user, financial_year):
    """Add a position to a financial year's register while the register is still being prepared."""
    if not register_editable(financial_year):
        raise DomainError(f"The {budget.fy_label(financial_year)} register is approved or in approval. "
                          "Request an ad-hoc position instead.")
    values = _default_reports_to(_clean(values))
    errors = validate(values)
    if errors:
        raise DomainError(" ".join(errors))
    register_for(financial_year, create=True)
    p = ApprovedPosition(position_number=next_reference("POS"), financial_year=financial_year, status="Planned",
                         **{k: values.get(k) for k in PLANNED_EDITABLE})
    db.session.add(p)
    db.session.flush()
    audit.record("create", "position", p.id, after=snapshot(p))
    return p


def update(position, values, user):
    """Edit a position (HR). Before register approval every field can change; after it, and while an ad-hoc
    request is in approval, everything except the department and cost centre."""
    if position.status in ("Abolished", "Rejected"):
        raise DomainError(f"A {position.status.lower()} position can't be edited.")
    planned = position.status in ("Planned", "Returned")
    allowed = PLANNED_EDITABLE if planned else HR_EDITABLE
    values = _clean(values)
    merged = {f: values.get(f, getattr(position, f)) if f in allowed else getattr(position, f)
              for f in PLANNED_EDITABLE}
    errors = validate(merged, position)
    if errors:
        raise DomainError(" ".join(errors))
    before = snapshot(position)
    for f in allowed:
        if f in values:
            setattr(position, f, values[f])
    b, a = audit.changes(before, snapshot(position))
    if a:
        audit.record("update", "position", position.id, before=b, after=a)


def delete_planned(position, user):
    if position.status != "Planned" or not register_editable(position.financial_year):
        raise DomainError("Only planned positions on a register that isn't approved can be deleted.")
    audit.record("delete", "position", position.id, before=snapshot(position))
    ApprovedPosition.query.filter_by(reports_to_id=position.id).update({"reports_to_id": None})
    db.session.delete(position)


def submit_register(financial_year, user):
    reg = register_for(financial_year)
    if reg is None or not ApprovedPosition.query.filter_by(financial_year=financial_year, status="Planned").count():
        raise DomainError(f"Add positions to the {budget.fy_label(financial_year)} register first.")
    if reg.status not in ("Draft", "Returned", "Rejected"):
        raise DomainError(f"The {budget.fy_label(financial_year)} register is already {reg.status.lower()}.")
    reg.status = "In Approval"
    reg.submitted_by_id = user.id
    reg.submitted_at = now()
    audit.record("status_change", "position_register", reg.id, after={"status": "In Approval",
                                                                       "fy": financial_year})
    return workflow.start("position_register", "position_register", reg.id, user,
                          {"financial_year": financial_year, "positions": planned_count(financial_year)})


def planned_count(financial_year):
    return ApprovedPosition.query.filter_by(financial_year=financial_year, status="Planned").count()


def _register(instance):
    return db.session.get(PositionRegister, instance.entity_id)


def _register_approved(instance):
    reg = _register(instance)
    reg.status = "Approved"
    reg.approved_at = now()
    for p in ApprovedPosition.query.filter_by(financial_year=reg.financial_year, status="Planned").all():
        set_status(p, "Filled" if p.holder else "Vacant", detail=f"{budget.fy_label(reg.financial_year)} register "
                                                                   "approved")
    audit.record("status_change", "position_register", reg.id, after={"status": "Approved"})


def _register_outcome(status):
    def handler(instance):
        reg = _register(instance)
        reg.status = status
        audit.record("status_change", "position_register", reg.id, after={"status": status})
    return handler


def _describe_register(instance):
    reg = _register(instance)
    label = budget.fy_label(reg.financial_year)
    return {"ref": f"Register {label}", "title": f"Approved positions register {label}",
            "link": f"/positions/?fy={reg.financial_year}"}


# ------------------------------------------------------------ ad-hoc requests

def ad_hoc_department(user):
    """The department a Department Manager may request ad-hoc positions for (their own); None for HR and others,
    who may choose any department."""
    if user.is_hr or not user.has_role(HOD) or not user.department_id:
        return None
    return db.session.get(Department, user.department_id)


def ad_hoc_reports_to(user):
    """The position a Department Manager's ad-hoc position reports to: their own (None if it isn't on the register)."""
    if ad_hoc_department(user) is None or user.employee is None:
        return None
    return user.employee.approved_position


def ad_hoc_defaults(user):
    """Pre-filled fields on a new ad-hoc request: a Department Manager's department, their cost centre (if it
    belongs to that department), the department's default job family and their own position as "reports to"."""
    dept = ad_hoc_department(user)
    if dept is None:
        return {}
    own_cc = user.employee.cost_centre_id if user.employee else None
    centres = [c.id for c in dept.cost_centres]
    manager_position = ad_hoc_reports_to(user)
    return {"department_id": dept.id, "cost_centre_id": own_cc if own_cc in centres else (centres[0] if centres else None),
            "job_family_id": dept.job_family_id, "reports_to_id": manager_position.id if manager_position else None}


def validate_ad_hoc(values):
    errors = validate(values)
    if len((values.get("justification") or "").strip()) < MIN_JUSTIFICATION:
        errors.append(f"The justification must be at least {MIN_JUSTIFICATION} characters.")
    return errors


def request_ad_hoc(values, user, position=None):
    """Create (or resubmit a returned) ad-hoc position request and start its approval."""
    values = _clean(values)
    values["justification"] = (values.get("justification") or "").strip()
    own = ad_hoc_department(user)
    if own is not None:
        if values.get("department_id") != own.id:
            raise DomainError(f"You can only request ad-hoc positions for {own.name}.")
        manager_position = ad_hoc_reports_to(user)            # always reports to the Department Manager
        values["reports_to_id"] = manager_position.id if manager_position else None
    else:
        values = _default_reports_to(values, position)
    errors = validate_ad_hoc(values)
    if errors:
        raise DomainError(" ".join(errors))
    fields = PLANNED_EDITABLE + ["justification"]
    if position is None:
        position = ApprovedPosition(position_number=next_reference("POS"), status="Pending approval", is_ad_hoc=True,
                                    financial_year=budget.financial_year_for(today()), requested_by_id=user.id,
                                    **{k: values.get(k) for k in fields})
        db.session.add(position)
        db.session.flush()
        audit.record("create", "position", position.id, after=snapshot(position))
    else:
        if position.status != "Returned" or position.requested_by_id != user.id:
            raise DomainError("Only the requester can resubmit a returned ad-hoc position.")
        before = snapshot(position)
        for k in fields:
            setattr(position, k, values.get(k))
        b, a = audit.changes(before, snapshot(position))
        audit.record("update", "position", position.id, before=b, after=a)
        set_status(position, "Pending approval", detail="Resubmitted")
    db.session.flush()
    return workflow.start("adhoc_position", "position", position.id, user,
                          {"department_id": position.department_id,
                           "budgeted_ctc": float(position.budgeted_ctc or 0),
                           "raised_by": "HR" if user.is_hr else "Manager"},
                          department=position.department)


def _position(instance):
    return db.session.get(ApprovedPosition, instance.entity_id)


def _adhoc_approved(instance):
    p = _position(instance)
    set_status(p, "Vacant", detail="Ad-hoc position approved and added to the register")
    if p.requested_by:
        notify(p.requested_by, "position_status", link=f"/positions/{p.id}", ref=p.position_number, title=p.title,
               status="Vacant", note="Your ad-hoc position is approved and on the register. You can now raise a "
                                     "requisition for it.")


def _adhoc_returned(instance):
    set_status(_position(instance), "Returned")


def _adhoc_rejected(instance):
    set_status(_position(instance), "Rejected")


def _describe_position(instance):
    p = _position(instance)
    return {"ref": p.position_number, "title": f"Ad-hoc position: {p.title}", "link": f"/positions/{p.id}"}


# --------------------------------------------- freeze / abolish / regrade

def pending_change(position):
    return PositionChange.query.filter_by(position_id=position.id, status="In Approval").first()


def request_change(position, kind, reason, user, new_job_grade_id=None):
    if kind not in CHANGE_FROM:
        raise DomainError("Choose freeze, unfreeze, abolish or regrade.")
    if position.status not in CHANGE_FROM[kind]:
        verb = {"Freeze": "frozen", "Unfreeze": "unfrozen", "Abolish": "abolished", "Regrade": "regraded"}[kind]
        raise DomainError(f"A {position.status.lower()} position can't be {verb}.")
    if pending_change(position):
        raise DomainError("This position already has a change waiting for approval.")
    if not (reason or "").strip():
        raise DomainError("Give a reason for the change.")
    if kind == "Regrade":
        grade = db.session.get(JobGrade, new_job_grade_id or 0)
        if grade is None or grade.id == position.job_grade_id:
            raise DomainError("Choose the new grade.")
    else:
        new_job_grade_id = None
    change = PositionChange(position_id=position.id, kind=kind, reason=reason.strip(), requested_by_id=user.id,
                            new_job_grade_id=new_job_grade_id, status="In Approval")
    db.session.add(change)
    db.session.flush()
    audit.record("create", "position_change", change.id,
                 after={"position": position.position_number, "kind": kind, "new_job_grade_id": new_job_grade_id})
    return workflow.start("position_change", "position_change", change.id, user,
                          {"kind": kind, "department_id": position.department_id}, department=position.department)


def _change(instance):
    return db.session.get(PositionChange, instance.entity_id)


def _change_approved(instance):
    change = _change(instance)
    change.status = "Approved"
    p = change.position
    if change.kind == "Regrade":
        before = p.job_grade_id
        p.job_grade_id = change.new_job_grade_id
        audit.record("update", "position", p.id, before={"job_grade_id": before},
                     after={"job_grade_id": p.job_grade_id}, detail="Regrade approved")
    else:
        target = {"Freeze": "Frozen", "Unfreeze": "Vacant", "Abolish": "Abolished"}[change.kind]
        if p.status in CHANGE_FROM[change.kind]:
            set_status(p, target, detail=f"{change.kind} approved")


def _change_outcome(status):
    def handler(instance):
        _change(instance).status = status
    return handler


def _describe_change(instance):
    change = _change(instance)
    p = change.position
    return {"ref": p.position_number, "title": f"{change.kind}: {p.title}", "link": f"/positions/{p.id}"}


workflow.register("position_register", describe=_describe_register, on_approved=_register_approved,
                  on_returned=_register_outcome("Returned"), on_rejected=_register_outcome("Rejected"),
                  sample_context={"financial_year": 2026, "positions": 10})
workflow.register("adhoc_position", describe=_describe_position, on_approved=_adhoc_approved,
                  on_returned=_adhoc_returned, on_rejected=_adhoc_rejected,
                  sample_context={"department_id": 1, "budgeted_ctc": 0.0, "raised_by": "Manager"})
workflow.register("position_change", describe=_describe_change, on_approved=_change_approved,
                  on_returned=_change_outcome("Returned"), on_rejected=_change_outcome("Rejected"),
                  sample_context={"kind": "Freeze", "department_id": 1})


# ---------------------------------------------------------------- CSV load

CSV_COLUMNS = ["financial_year", "department_code", "cost_centre_code", "title", "grade", "branch_code",
               "job_family_code", "budgeted_ctc", "reports_to", "minimum_requirements"]


def csv_template():
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(CSV_COLUMNS)
    fy = budget.fy_label(budget.financial_year_for(today()))[2:]
    cc = CostCentre.query.join(Department).order_by(Department.code, CostCentre.code).first()
    if cc is not None:
        w.writerow([fy, cc.department.code, cc.code, "Sales Consultant", "Junior", "", "", "250000.00", "",
                    "Matric | 1 year of sales experience | Valid driver's licence"])
    return out.getvalue()


def import_csv(text, user):
    """All-or-nothing load of planned positions. Returns (rows_saved, errors).
    reports_to is an existing position number; minimum requirements are separated by '|'."""
    reader = csv.DictReader(io.StringIO(text.lstrip("﻿")))
    header = [h.strip().lower() for h in (reader.fieldnames or [])]
    if header != CSV_COLUMNS:
        return 0, [f"Row 1: the header must be {', '.join(CSV_COLUMNS)}."]
    departments = {d.code.upper(): d for d in Department.query.all()}
    centres = {c.code.upper(): c for c in CostCentre.query.all()}
    grades = {g.name.lower(): g for g in JobGrade.query.all()}            # Intern, Junior, Senior, Manager, Exco
    branches = {b.code.upper(): b for b in Branch.query.all()}
    families = {f.code.upper(): f for f in JobFamily.query.all()}
    numbers = {p.position_number.upper(): p for p in ApprovedPosition.query.all()}
    errors, parsed = [], []
    for n, raw in enumerate(reader, start=2):
        row = {k.strip().lower(): (v or "").strip() for k, v in raw.items() if k}
        problems = []
        try:
            fy = budget.parse_fy(row["financial_year"])
            if not register_editable(fy):
                problems.append(f"the {budget.fy_label(fy)} register is approved or in approval")
        except ValueError as exc:
            fy = None
            problems.append(str(exc))
        dept = departments.get(row["department_code"].upper())
        cc = centres.get(row["cost_centre_code"].upper())
        if dept is None:
            problems.append(f"unknown department '{row['department_code']}'")
        if cc is None:
            problems.append(f"unknown cost centre '{row['cost_centre_code']}'")
        elif dept is not None and cc.department_id != dept.id:
            problems.append(f"cost centre {cc.code} doesn't belong to {dept.code}")
        if not row["title"]:
            problems.append("title is required")
        grade = grades.get(row["grade"].lower())
        if grade is None:
            problems.append(f"unknown grade '{row['grade']}' (use Intern, Junior, Senior, Manager or Exco)")
        branch = branches.get(row["branch_code"].upper()) if row["branch_code"] else None
        if row["branch_code"] and branch is None:
            problems.append(f"unknown branch '{row['branch_code']}'")
        family = families.get(row["job_family_code"].upper()) if row["job_family_code"] else None
        if row["job_family_code"] and family is None:
            problems.append(f"unknown job family '{row['job_family_code']}'")
        ctc = None
        if row["budgeted_ctc"]:
            try:
                ctc = Decimal(row["budgeted_ctc"].replace(" ", "").replace(",", ""))
                if not math.isfinite(ctc) or ctc < 0:
                    raise InvalidOperation
                ctc = ctc.quantize(Decimal("0.01"))
            except (InvalidOperation, ValueError):
                problems.append(f"budgeted CTC '{row['budgeted_ctc']}' isn't a valid number")
        reports_to = numbers.get(row["reports_to"].upper()) if row["reports_to"] else None
        if row["reports_to"] and reports_to is None:
            problems.append(f"unknown reports-to position '{row['reports_to']}'")
        if problems:
            errors.append(f"Row {n}: " + "; ".join(problems) + ".")
            continue
        requirements = "\n".join(r.strip() for r in row["minimum_requirements"].split("|") if r.strip()) or None
        parsed.append((fy, {"title": row["title"], "department_id": dept.id, "cost_centre_id": cc.id,
                            "job_grade_id": grade.id, "branch_id": branch.id if branch else None,
                            "job_family_id": family.id if family else None, "budgeted_ctc": ctc,
                            "reports_to_id": reports_to.id if reports_to else None,
                            "minimum_requirements": requirements, "critical": False}))
    if not parsed and not errors:
        errors.append("The file has no data rows.")
    if errors:
        return 0, errors
    for fy, values in parsed:
        create_planned(values, user, fy)
    audit.record("position.import", "position", None, after={"rows": len(parsed)})
    return len(parsed), []


# --------------------------------------------------------- vacancy dashboard

GROUPS = {
    "business_unit": ("Business unit", lambda p: p.business_unit),
    "department": ("Department", lambda p: p.department),
    "branch": ("Branch / area", lambda p: p.branch),
    "region": ("Region", lambda p: p.region),
    "cost_centre": ("Cost centre", lambda p: p.cost_centre),
    "grade": ("Grade", lambda p: p.job_grade),
    "job_family": ("Job family", lambda p: p.job_family),
}


def _bucket(position, req_status):
    """Which vacancy column a position on the register counts in."""
    if position.status == "Filled":
        return "filled"
    if position.status == "Frozen":
        return "frozen"
    if position.status == "In recruitment" and req_status in ADVERTISED_REQUISITION:
        return "in_recruitment"
    return "not_advertised"


def _counts(positions, req_status):
    c = {"approved": len(positions), "filled": 0, "frozen": 0, "in_recruitment": 0, "not_advertised": 0}
    for p in positions:
        c[_bucket(p, req_status.get(p.id))] += 1
    c["vacant"] = c["frozen"] + c["in_recruitment"] + c["not_advertised"]
    c["rate"] = round(c["vacant"] * 100 / c["approved"]) if c["approved"] else 0
    return c


def vacancy_dashboard(user, financial_year, group_by="business_unit"):
    group_by = group_by if group_by in GROUPS else "business_unit"
    positions = (visible_query(user).filter(ApprovedPosition.financial_year == financial_year,
                                            ApprovedPosition.status.in_(ON_REGISTER)).all())
    req_status = dict(db.session.query(Requisition.approved_position_id, Requisition.status)
                      .filter(Requisition.approved_position_id.in_([p.id for p in positions] or [0]),
                              Requisition.status.in_(OPEN_REQUISITION)).all())
    totals = _counts(positions, req_status)
    vacant = [p for p in positions if p.status in VACANT]
    days = [(today() - p.vacant_since).days for p in vacant if p.vacant_since]
    filled = [p.days_to_fill for p in positions if p.status == "Filled" and p.days_to_fill is not None]
    totals["avg_days_vacant"] = round(sum(days) / len(days)) if days else 0
    totals["time_to_fill"] = round(sum(filled) / len(filled)) if filled else None
    totals["ad_hoc"] = sum(1 for p in positions if p.is_ad_hoc)

    label, key = GROUPS[group_by]
    grouped = {}
    for p in positions:
        grouped.setdefault(key(p), []).append(p)
    rows = sorted(({"group": g, "name": str(g) if g is not None else "Not set", **_counts(ps, req_status)}
                   for g, ps in grouped.items()), key=lambda r: (r["group"] is None, -r["approved"], r["name"]))

    longest = sorted((p for p in vacant if p.vacant_since and p.status != "Frozen"),
                     key=lambda p: p.vacant_since)[:5]
    year_start, year_end = budget.fy_bounds(financial_year)
    ad_hoc = visible_query(user).filter(ApprovedPosition.is_ad_hoc.is_(True),
                                        ApprovedPosition.created_at >= year_start,
                                        ApprovedPosition.created_at <= date.fromordinal(year_end.toordinal() + 1)).all()
    ad_hoc_summary = {"approved": sum(1 for p in ad_hoc if p.status in ON_REGISTER + ("Abolished",)),
                      "pending": [p for p in ad_hoc if p.status == "Pending approval"]}
    return {"totals": totals, "rows": rows, "group_label": label, "group_by": group_by, "longest": longest,
            "ad_hoc": ad_hoc_summary, "positions": positions, "req_status": req_status}


def dashboard_workbook(data, financial_year):
    """The vacancy dashboard as an Excel workbook: a summary sheet and every position on the register."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Vacancies"
    ws.append([f"Vacancy dashboard {budget.fy_label(financial_year)}, by {data['group_label'].lower()}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append([data["group_label"], "Approved", "Filled", "Vacant", "In recruitment", "Not yet advertised",
               "Frozen", "Vacancy rate %"])
    for cell in ws[3]:
        cell.font = Font(bold=True)
    for r in data["rows"] + [{"name": "Icebolethu Group", **data["totals"]}]:
        ws.append([r["name"], r["approved"], r["filled"], r["vacant"], r["in_recruitment"], r["not_advertised"],
                   r["frozen"], r["rate"]])
    ws[ws.max_row][0].font = Font(bold=True)

    detail = wb.create_sheet("Positions")
    detail.append(["Position", "Title", "Business unit", "Department", "Branch / area", "Region", "Cost centre",
                   "Grade", "Job family", "Status", "Ad hoc", "Vacant since", "Days vacant", "Holder"])
    for cell in detail[1]:
        cell.font = Font(bold=True)
    for p in sorted(data["positions"], key=lambda p: p.position_number):
        detail.append([p.position_number, p.title, str(p.business_unit or ""), p.department.name,
                       str(p.branch or ""), str(p.region or ""), p.cost_centre.code if p.cost_centre else "",
                       p.job_grade.name if p.job_grade else "", str(p.job_family or ""), p.status,
                       "Yes" if p.is_ad_hoc else "", p.vacant_since,
                       (today() - p.vacant_since).days if p.vacant_since else "",
                       p.holder.full_name if p.holder else ""])
    for sheet in (ws, detail):
        for col in sheet.columns:
            width = max(len(str(c.value or "")) for c in col)
            sheet.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 45)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def reference_data():
    departments = Department.query.filter_by(active=True).order_by(Department.name).all()
    managers = {d.id: department_manager_position(d.id) for d in departments}
    return {
        "departments": departments,
        "manager_positions": {d_id: p.id for d_id, p in managers.items() if p is not None},
        "cost_centres": CostCentre.query.filter_by(active=True).order_by(CostCentre.code).all(),
        "grades": JobGrade.query.order_by(JobGrade.grade).all(),
        "branches": Branch.query.filter_by(active=True).order_by(Branch.name).all(),
        "families": JobFamily.query.filter_by(active=True).order_by(JobFamily.name).all(),
    }
