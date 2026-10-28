"""ATS engine: extracts text from CVs and scores candidates against a job.

The score is a weighted percentage made up of:
  - required skills matched      (45%)
  - preferred skills matched     (15%)
  - years of experience          (20%)
  - education level              (10%)
  - job-description keywords     (10%)
Components that don't apply to a job (e.g. no preferred skills listed) are
dropped and the remaining weights are rescaled so the total is still 100%.
"""
import io
import logging
import re
from datetime import date

log = logging.getLogger(__name__)

WEIGHTS = {
    "required": 45,
    "preferred": 15,
    "experience": 20,
    "education": 10,
    "keywords": 10,
}

EDUCATION_LEVELS = [
    (0, "None specified"),
    (1, "Matric / Grade 12"),
    (2, "Certificate"),
    (3, "Diploma"),
    (4, "Bachelor's Degree"),
    (5, "Honours / Postgraduate Diploma"),
    (6, "Master's Degree"),
    (7, "Doctorate / PhD"),
]
EDUCATION_LABELS = dict(EDUCATION_LEVELS)

# Highest-level patterns first so the first hit is the best qualification.
_EDU_PATTERNS = [
    (7, r"\bph\.?\s?d\b|\bdoctorate\b|\bdoctor of\b"),
    (6, r"\bmaster'?s?\b|\bmba\b|\bm\.?sc\b|\bm\.?com\b|\bm\.?eng\b"),
    (5, r"\bhonours\b|\bhons\b|\bpostgraduate diploma\b|\bpgdip\b|\bb\.?tech\b"),
    (4, r"\bbachelor'?s?\b|\bdegree\b|\bb\.?com\b|\bb\.?sc\b|\bb\.?a\.?\b|\bllb\b|\bb\.?eng\b|\bbba\b"),
    (3, r"\bdiploma\b|\bn6\b|\bnational diploma\b"),
    (2, r"\bcertificate\b|\bnqf\s?[4-5]\b|\bcertified\b"),
    (1, r"\bmatric\b|\bgrade\s?12\b|\bnsc\b|\bnational senior certificate\b"),
]

_STOPWORDS = set("""
a about above after again all also an and any are as at be because been before
being below between both but by can could did do does doing down during each
few for from further had has have having he her here hers him his how i if in
into is it its itself just me more most my no nor not of off on once only or
other our ours out over own same she should so some such than that the their
them then there these they this those through to too under until up very was
we were what when where which while who whom why will with would you your
able ability work working role position candidate candidates job team company
must strong excellent good great well within including include includes etc
responsible responsibilities requirements required preferred minimum years year
experience knowledge skills skill using use new per key ensure provide support
""".split())


# ---------------------------------------------------------------- extraction

REQUIRED_PACKAGES = {"pypdf": "pypdf", "docx": "python-docx"}  # import name -> pip name


def missing_packages():
    """pip names of CV-reading libraries that can't be imported in this Python."""
    missing = []
    for module, pip_name in REQUIRED_PACKAGES.items():
        try:
            __import__(module)
        except ImportError:
            missing.append(pip_name)
    return missing


def extract_text(filename, data):
    """Return plain text from a PDF, DOCX or TXT file's bytes.

    A missing library raises ImportError on purpose: it must never be
    mistaken for a CV that has no readable text."""
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        try:
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                reader.decrypt("")  # many "protected" CVs only restrict editing, not opening
            return "\n".join((page.extract_text() or "") for page in reader.pages)
        except Exception as exc:
            log.warning("Could not read PDF %s: %s", filename, exc)
            return ""
    if name.endswith(".docx"):
        import docx
        try:
            document = docx.Document(io.BytesIO(data))
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(parts)
        except Exception as exc:
            log.warning("Could not read DOCX %s: %s", filename, exc)
            return ""
    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore")
    return ""


# ------------------------------------------------------------------ helpers

def split_skills(value):
    """'Excel, SQL; Power BI' -> ['excel', 'sql', 'power bi']"""
    if not value:
        return []
    seen, out = set(), []
    for part in re.split(r"[,;\n]", value):
        skill = part.strip().lower()
        if skill and skill not in seen:
            seen.add(skill)
            out.append(skill)
    return out


def _normalise(text):
    text = (text or "").lower().replace("’", "").replace("'", "")
    return re.sub(r"\s+", " ", text)


def has_term(text, term):
    """Whole-term match that also works for terms like 'c++' or 'node.js',
    and accepts plurals ('specification' matches 'specifications')."""
    plural = "(?:s|es)?" if term[-1:].isalpha() else ""
    pattern = (r"(?<![a-z0-9])" + re.escape(term).replace(r"\ ", r"[\s\-_/]*")
               + plural + r"(?![a-z0-9])")
    return re.search(pattern, text) is not None


def _variants(skill):
    """Common spellings of a skill, e.g. 'ms excel' also matches 'excel'."""
    skill = _normalise(skill)
    variants = {skill, re.sub(r"^(ms|microsoft)\s+", "", skill)}
    for v in list(variants):
        if v.endswith("analysis"):
            variants.add(v[:-len("analysis")] + "analyst")
        variants.add(v.replace("modelling", "modeling"))
        variants.add(v.replace("licence", "license"))
        variants.add(v.replace("license", "licence"))
    return [v for v in variants if v]


def has_skill(text, skill):
    return any(has_term(text, v) for v in _variants(skill))


_MONTHS = {m: i for i, m in enumerate(
    "jan feb mar apr may jun jul aug sep oct nov dec".split(), 1)}
_MON = r"(?:(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?|(\d{1,2})\s*/)\s*"
_RANGE = re.compile(
    rf"(?:{_MON})?((?:19|20)\d\d)\s*(?:-|–|—|to|until)\s*"
    rf"(?:(?:{_MON})?((?:19|20)\d\d)|(present|current|now|date|today))", re.I)
# Lines about studies, so the years spent studying aren't counted as work.
_EDU_LINE = re.compile(r"universit|college|school|academy|institut|tvet|campus|diploma|degree|"
                       r"bachelor|honours|matric|grade\s?12|certificate|studies|qualification", re.I)


def _month_index(name, number, year, default):
    if name:
        month = _MONTHS[name[:3].lower()]
    elif number and 1 <= int(number) <= 12:
        month = int(number)
    else:
        month = default
    return int(year) * 12 + month - 1


def experience_from_dates(text, today=None):
    """Years of work experience from employment date ranges such as
    'Apr 2023 – Apr 2025', '2019 - 2022' or '03/2020 – Present'.
    Overlapping jobs are merged and study periods are skipped."""
    today = today or date.today()
    now = today.year * 12 + today.month - 1
    spans, previous = [], ""
    for line in (text or "").splitlines():
        line = line.strip()
        if not line:
            continue
        if not (_EDU_LINE.search(line) or _EDU_LINE.search(previous)):
            for m in _RANGE.finditer(line):
                start = _month_index(m.group(1), m.group(2), m.group(3), 1)
                end = now if m.group(7) else _month_index(m.group(4), m.group(5), m.group(6), 12)
                end = min(end, now)
                if 0 <= end - start <= 45 * 12:
                    spans.append((start, end + 1))
        previous = line
    months, current_end = 0, None
    for start, end in sorted(spans):
        if current_end is None or start >= current_end:
            months += end - start
            current_end = end
        elif end > current_end:
            months += end - current_end
            current_end = end
    return round(months / 12, 1)


def detect_years(text):
    """Best guess at total years of experience in the CV: an explicit
    'N years of experience' phrase, or the sum of employment date ranges."""
    lowered = (text or "").lower()
    found = [int(n) for n in re.findall(
        r"(\d{1,2})\s*\+?\s*(?:years?|yrs?)(?:\s+of)?(?:\s+\w+){0,3}?\s+experience", lowered)]
    if not found:
        found = [int(n) for n in re.findall(r"(\d{1,2})\s*\+?\s*(?:years?|yrs?)\b", lowered)]
    found = [n for n in found if n <= 50]
    return max([float(n) for n in found] + [experience_from_dates(text)])


def detect_education(text):
    for level, pattern in _EDU_PATTERNS:
        if re.search(pattern, text):
            return level
    return 0


def _keywords(text, limit=40):
    counts = {}
    for word in re.findall(r"[a-z][a-z+#.\-]{2,}", text):
        word = word.strip(".-")
        if len(word) < 3 or word in _STOPWORDS:
            continue
        counts[word] = counts.get(word, 0) + 1
    ranked = sorted(counts.items(), key=lambda kv: (-kv[1], kv[0]))
    return [w for w, _ in ranked[:limit]]


# ------------------------------------------------------------------ scoring

def score_candidate(job, candidate):
    """Score a candidate against a job.

    `job` and `candidate` are mappings (e.g. sqlite3.Row) with the columns
    used below. Returns (score_percent, breakdown_dict).
    """
    cv_text = _normalise(candidate["cv_text"])
    profile = _normalise(" ".join([
        cv_text,
        candidate["skills"] or "",
        candidate["current_title"] or "",
    ]))

    components = {}

    required = split_skills(job["required_skills"])
    if required:
        matched = [s for s in required if has_skill(profile, s)]
        components["required"] = {
            "score": 100 * len(matched) / len(required),
            "matched": matched,
            "missing": [s for s in required if s not in matched],
        }

    preferred = split_skills(job["preferred_skills"])
    if preferred:
        matched = [s for s in preferred if has_skill(profile, s)]
        components["preferred"] = {
            "score": 100 * len(matched) / len(preferred),
            "matched": matched,
            "missing": [s for s in preferred if s not in matched],
        }

    min_exp = job["min_experience"] or 0
    stated = candidate["years_experience"]
    # Use the higher of the applicant's answer and what the CV shows, so a
    # "0" typed on the form can't hide real experience listed in the CV.
    years = max(stated or 0, detect_years(candidate["cv_text"]))
    if min_exp > 0:
        components["experience"] = {
            "score": min(1.0, years / min_exp) * 100,
            "candidate_years": years,
            "required_years": min_exp,
        }

    min_edu = job["min_education"] or 0
    edu = max(candidate["education_level"] or 0, detect_education(cv_text))
    if min_edu > 0:
        components["education"] = {
            "score": 100 if edu >= min_edu else 100 * edu / min_edu,
            "candidate_level": EDUCATION_LABELS.get(edu, "Unknown"),
            "required_level": EDUCATION_LABELS.get(min_edu, "Unknown"),
        }

    keywords = _keywords(_normalise(f"{job['title']} {job['description']}"))
    if keywords:
        hits = [k for k in keywords if has_term(profile, k)]
        # Hitting ~60% of the job's vocabulary is treated as a full match.
        components["keywords"] = {
            "score": min(100.0, 100 * len(hits) / (len(keywords) * 0.6)),
            "matched": hits[:15],
        }

    if not components:
        return 0.0, {"components": {}, "note": "Job has no screening criteria."}

    total_weight = sum(WEIGHTS[k] for k in components)
    score = sum(WEIGHTS[k] * c["score"] for k, c in components.items()) / total_weight
    for key, comp in components.items():
        comp["weight"] = round(100 * WEIGHTS[key] / total_weight, 1)
        comp["score"] = round(comp["score"], 1)

    breakdown = {
        "components": components,
        "detected_years": years,
        "detected_education": EDUCATION_LABELS.get(edu, "Unknown"),
        "cv_text_found": bool(cv_text.strip()),
    }
    return round(score, 1), breakdown
